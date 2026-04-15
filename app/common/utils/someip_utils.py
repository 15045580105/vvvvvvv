import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import openpyxl

from app.common.utils.path_utils import get_device_config_path
from app.services.system.config_service import ConfigService


class SomeIPMatrixError(Exception):
    """SomeIP matrix parsing error."""


class SomeIPMatrixStore:
    """Parse and cache SomeIP matrix data.

    Features:
    - Deployment sheet: service_id + instance_id -> service interface name
    - Payload sheets: service_id + method_id -> payload field rules
    """

    CACHE_NAME = "someip_matrix_cache"
    DEPLOYMENT_SHEET_NAME = "Deployment"
    SERVICE_INTERFACES_SHEET_NAME = "ServiceInterfaces"
    DATA_TYPE_DEFINITION_SHEET_NAME = "DataTypeDefinition"
    SERIALIZATION_PARAMETER_SHEET_NAME = "SerializationParameter"
    PAYLOAD_HEADER_SCAN_LIMIT = 40

    def __init__(self) -> None:
        self._config_service = ConfigService()
        self._memory_cache: Dict[str, Dict[str, Any]] = {}

    def import_matrix(self, file_path: Optional[str] = None, force_reload: bool = False) -> Dict[str, Any]:
        matrix_path = self.resolve_matrix_path(file_path)
        fingerprint = self._build_fingerprint(matrix_path)
        memory_key = str(matrix_path)

        if not force_reload:
            cached = self._memory_cache.get(memory_key)
            if cached and cached.get("fingerprint") == fingerprint:
                return cached

            disk_cache = self._load_disk_cache()
            cache_entry = disk_cache.get(memory_key)
            if cache_entry and cache_entry.get("fingerprint") == fingerprint:
                self._memory_cache[memory_key] = cache_entry
                return cache_entry

        parsed = self._parse_matrix_file(matrix_path)
        payload = {
            "matrix_path": str(matrix_path),
            "fingerprint": fingerprint,
            "services": parsed.get("services", {}),
            "payload_rules": {},
            "message_schemas": {},
            "datatype_registry": {},
            "serialization_config": {},
            "payload_rules_loaded": False,
            "count": len(parsed.get("services", {})),
            "payload_rule_count": 0,
        }
        self._memory_cache[memory_key] = payload

        disk_cache = self._load_disk_cache()
        disk_cache[memory_key] = payload
        self._save_disk_cache(disk_cache)
        return payload

    def get_service_interface_name(
        self,
        service_id: Any,
        instance_id: Any,
        file_path: Optional[str] = None,
        default: str = "Unknown Service",
    ) -> str:
        try:
            cache = self.import_matrix(file_path=file_path)
        except Exception:
            return default
        lookup_key = self._build_lookup_key(service_id, instance_id)
        item = cache.get("services", {}).get(lookup_key)
        if item:
            return item.get("service_interface_name") or default
        return default

    def get_service_info(
        self,
        service_id: Any,
        instance_id: Any,
        file_path: Optional[str] = None,
    ) -> Optional[Dict[str, Any]]:
        try:
            cache = self.import_matrix(file_path=file_path)
        except Exception:
            return None
        return cache.get("services", {}).get(self._build_lookup_key(service_id, instance_id))

    def get_payload_rules(
        self,
        service_id: Any,
        method_id: Any,
        file_path: Optional[str] = None,
        message_type: Optional[Any] = None,
    ) -> List[Dict[str, Any]]:
        try:
            cache = self.import_matrix(file_path=file_path)
        except Exception:
            return []

        if isinstance(cache, dict) and not cache.get("payload_rules_loaded"):
            matrix_path = Path(str(cache.get("matrix_path") or self.resolve_matrix_path(file_path)))
            rules = self._load_payload_rules(matrix_path, cache)
        else:
            rules = cache.get("payload_rules", {}) if isinstance(cache, dict) else {}

        normalized_service_id = self._normalize_id(service_id)
        normalized_method_id = self._normalize_id(method_id)
        message_kind = self._normalize_message_kind(message_type)

        candidate_keys: List[str] = []
        if message_kind:
            candidate_keys.extend([
                f"{normalized_service_id}:{normalized_method_id}:{message_kind}",
                f"{normalized_service_id}:*:{message_kind}",
            ])
        candidate_keys.extend([
            f"{normalized_service_id}:{normalized_method_id}",
            f"{normalized_service_id}:*",
        ])

        for key in candidate_keys:
            if key in rules:
                return list(rules.get(key) or [])
        return []

    def resolve_matrix_path(self, file_path: Optional[str] = None) -> Path:
        candidates: List[Path] = []
        if file_path:
            candidates.append(Path(file_path))

        settings = self._load_app_settings().get("someip", {})
        configured = settings.get("matrix_file")
        if configured:
            candidates.append(Path(configured))

        utils_dir = Path(__file__).resolve().parent
        common_file_dir = utils_dir.parent / "file"
        candidates.extend(sorted(common_file_dir.glob("*.xlsx")))
        candidates.extend(sorted(common_file_dir.glob("*.xlsm")))
        candidates.extend(sorted(utils_dir.glob("*.xlsx")))
        candidates.extend(sorted(utils_dir.glob("*.xlsm")))

        for candidate in candidates:
            if candidate and candidate.exists() and candidate.is_file():
                return candidate.resolve()

        raise FileNotFoundError("未找到 SomeIP 矩阵表文件，请先在 SomeIP 设置界面中导入矩阵 Excel 文件。")

    def _parse_matrix_file(self, file_path: Path) -> Dict[str, Any]:
        workbook = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        try:
            if self.DEPLOYMENT_SHEET_NAME not in workbook.sheetnames:
                raise SomeIPMatrixError(f"矩阵缺少 `{self.DEPLOYMENT_SHEET_NAME}` sheet")

            services = self._parse_deployment_sheet(workbook[self.DEPLOYMENT_SHEET_NAME])
            return {
                "services": services,
            }
        finally:
            workbook.close()

    def _load_payload_rules(self, matrix_path: Path, cache_entry: Dict[str, Any]) -> Dict[str, List[Dict[str, Any]]]:
        workbook = openpyxl.load_workbook(str(matrix_path), read_only=True, data_only=True)
        try:
            services = cache_entry.get("services", {}) if isinstance(cache_entry, dict) else {}
            parsed_payload = self._parse_payload_metadata(workbook, services)
        finally:
            workbook.close()

        payload_rules = parsed_payload.get("message_schemas", {})
        cache_entry["payload_rules"] = payload_rules
        cache_entry["message_schemas"] = payload_rules
        cache_entry["datatype_registry"] = parsed_payload.get("datatype_registry", {})
        cache_entry["serialization_config"] = parsed_payload.get("serialization_config", {})
        cache_entry["payload_rules_loaded"] = True
        cache_entry["payload_rule_count"] = sum(len(items) for items in payload_rules.values())

        memory_key = str(matrix_path)
        self._memory_cache[memory_key] = cache_entry
        disk_cache = self._load_disk_cache()
        disk_cache[memory_key] = cache_entry
        self._save_disk_cache(disk_cache)
        return payload_rules

    def _parse_payload_metadata(self, workbook, services: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
        if (
            self.SERVICE_INTERFACES_SHEET_NAME in workbook.sheetnames
            and self.DATA_TYPE_DEFINITION_SHEET_NAME in workbook.sheetnames
        ):
            service_entries = self._parse_service_interfaces_sheet(workbook[self.SERVICE_INTERFACES_SHEET_NAME])
            datatype_registry = self._parse_datatype_definition_sheet(workbook[self.DATA_TYPE_DEFINITION_SHEET_NAME])
            serialization_config = self._parse_serialization_sheet(
                workbook[self.SERIALIZATION_PARAMETER_SHEET_NAME]
            ) if self.SERIALIZATION_PARAMETER_SHEET_NAME in workbook.sheetnames else self._default_serialization_config()
            message_schemas = self._build_message_schemas(service_entries, datatype_registry, serialization_config)
            return {
                "message_schemas": message_schemas,
                "datatype_registry": datatype_registry,
                "serialization_config": serialization_config,
            }

        return {
            "message_schemas": self._parse_payload_sheets(workbook, services),
            "datatype_registry": {},
            "serialization_config": self._default_serialization_config(),
        }

    def _parse_service_interfaces_sheet(self, sheet) -> List[Dict[str, Any]]:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        header_row_index = self._find_service_interface_header_row_index(rows)
        headers = [self._normalize_header(v) for v in rows[header_row_index]]

        service_name_idx = self._find_col_idx(headers, ["service interface name", "service name", "interface name"])
        service_id_idx = self._find_col_idx(headers, ["service id", "serviceid"])
        method_kind_idx = self._find_col_idx(headers, ["method/event/field", "method event field"])
        accessor_idx = self._find_col_idx(headers, ["setter/getter/notifier", "setter getter notifier"])
        element_name_idx = self._find_col_idx(headers, ["element name", "method name", "event name"])
        method_id_idx = self._find_col_idx(headers, ["method id/event id", "method id", "event id", "message id"])
        parameter_name_idx = self._find_col_idx(headers, ["parameter name", "parameter", "field name"])
        direction_idx = self._find_col_idx(headers, ["in/out", "direction"])
        parameter_type_idx = self._find_col_idx(headers, ["parameter data type", "parameter datatype", "data type"])
        parameter_desc_idx = self._find_col_idx(headers, ["parameter description"])
        if parameter_desc_idx < 0:
            parameter_desc_idx = self._find_col_idx(headers, ["param description"])
        element_desc_idx = self._find_col_idx(headers, ["element description"])

        if min(service_name_idx, service_id_idx, element_name_idx, method_id_idx, parameter_type_idx) < 0:
            return []

        entries: List[Dict[str, Any]] = []
        last_service_name = ""
        last_service_id = ""
        last_method_kind = ""
        last_accessor = ""
        last_element_name = ""
        last_method_id = ""

        for row_number, row in enumerate(rows[header_row_index + 1:], start=header_row_index + 2):
            service_name = self._safe_row_get(row, service_name_idx) or last_service_name
            service_id = self._normalize_id(self._safe_row_get(row, service_id_idx)) or last_service_id
            method_kind = self._safe_row_get(row, method_kind_idx) if method_kind_idx >= 0 else ""
            method_kind = method_kind or last_method_kind
            accessor = self._safe_row_get(row, accessor_idx) if accessor_idx >= 0 else ""
            accessor = accessor or last_accessor
            element_name = self._safe_row_get(row, element_name_idx) or last_element_name
            method_id = self._normalize_id(self._safe_row_get(row, method_id_idx)) or last_method_id

            if service_name:
                last_service_name = service_name
            if service_id:
                last_service_id = service_id
            if method_kind:
                last_method_kind = method_kind
            if accessor:
                last_accessor = accessor
            if element_name:
                last_element_name = element_name
            if method_id:
                last_method_id = method_id

            parameter_name = self._safe_row_get(row, parameter_name_idx)
            parameter_type = self._safe_row_get(row, parameter_type_idx)
            direction = self._safe_row_get(row, direction_idx).upper()
            element_description = self._safe_row_get(row, element_desc_idx) if element_desc_idx >= 0 else ""
            parameter_description = self._safe_row_get(row, parameter_desc_idx) if parameter_desc_idx >= 0 else ""
            normalized_method_kind = self._normalize_header(method_kind)
            normalized_accessor = self._normalize_header(accessor)

            is_field_row = "field" in normalized_method_kind and normalized_accessor in {"getter", "notifier"}
            if is_field_row and not parameter_name:
                parameter_name = element_name
            if is_field_row and not parameter_description:
                parameter_description = element_description

            if not service_id or not method_id or not parameter_name or not parameter_type:
                continue

            entries.append({
                "row": row_number,
                "service_interface_name": service_name,
                "service_id": service_id,
                "method_kind": method_kind,
                "accessor": accessor,
                "method_name": element_name,
                "method_id": method_id,
                "parameter_name": parameter_name,
                "direction": direction,
                "parameter_description": parameter_description,
                "parameter_data_type": parameter_type,
                "message_kind": self._resolve_message_kind_from_row(method_kind, accessor, direction),
                "is_field_row": is_field_row,
                "element_description": element_description,
            })

        return entries

    def _parse_datatype_definition_sheet(self, sheet) -> Dict[str, Dict[str, Any]]:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return {}

        header_row_index = self._find_datatype_header_row_index(rows)
        headers = [self._normalize_header(v) for v in rows[header_row_index]]

        name_idx = self._find_col_idx(headers, ["parameter data type name", "datatype name", "data type name"])
        description_idx = self._find_col_idx(headers, ["datatype description", "description"])
        category_idx = self._find_col_idx(headers, ["data category", "category"])
        length_type_idx = self._find_col_idx(headers, ["string/array length type", "length type"])
        length_min_idx = self._find_col_idx(headers, ["string/array length min", "length min"])
        length_max_idx = self._find_col_idx(headers, ["string/array length max", "length max"])
        member_name_idx = self._find_col_idx(headers, ["member name"])
        member_description_idx = self._find_col_idx(headers, ["member description"])
        member_reference_idx = self._find_col_idx(headers, ["member datatype reference", "datatype reference"])
        datatype_idx = self._find_col_idx(headers, ["datatype", "type"])
        resolution_idx = self._find_col_idx(headers, ["resolution", "factor", "scale"])
        offset_idx = self._find_col_idx(headers, ["offset"])
        unit_idx = self._find_col_idx(headers, ["unit"])
        discrete_idx = self._find_col_idx(headers, ["discrete value defination", "discrete value definition", "enum"])
        comment_idx = self._find_col_idx(headers, ["comment"])
        initial_value_idx = self._find_col_idx(headers, ["initial value"])
        invalid_value_idx = self._find_col_idx(headers, ["invalid value"])

        if name_idx < 0:
            return {}

        registry: Dict[str, Dict[str, Any]] = {}
        for row_number, row in enumerate(rows[header_row_index + 1:], start=header_row_index + 2):
            name = self._safe_row_get(row, name_idx)
            if not name:
                continue

            item = registry.setdefault(name, {
                "type_name": name,
                "description": self._safe_row_get(row, description_idx) if description_idx >= 0 else "",
                "data_category": self._safe_row_get(row, category_idx) if category_idx >= 0 else "",
                "length_type": self._safe_row_get(row, length_type_idx) if length_type_idx >= 0 else "",
                "length_min": self._safe_row_get(row, length_min_idx) if length_min_idx >= 0 else "",
                "length_max": self._safe_row_get(row, length_max_idx) if length_max_idx >= 0 else "",
                "data_type": self._safe_row_get(row, datatype_idx) if datatype_idx >= 0 else "",
                "resolution": self._safe_row_get(row, resolution_idx) if resolution_idx >= 0 else "",
                "offset": self._safe_row_get(row, offset_idx) if offset_idx >= 0 else "",
                "unit": self._safe_row_get(row, unit_idx) if unit_idx >= 0 else "",
                "discrete_values": self._safe_row_get(row, discrete_idx) if discrete_idx >= 0 else "",
                "comment": self._safe_row_get(row, comment_idx) if comment_idx >= 0 else "",
                "initial_value": self._safe_row_get(row, initial_value_idx) if initial_value_idx >= 0 else "",
                "invalid_value": self._safe_row_get(row, invalid_value_idx) if invalid_value_idx >= 0 else "",
                "members": [],
                "rows": [],
            })
            item["rows"].append(row_number)

            if not item.get("description") and description_idx >= 0:
                item["description"] = self._safe_row_get(row, description_idx)
            if not item.get("data_category") and category_idx >= 0:
                item["data_category"] = self._safe_row_get(row, category_idx)
            if not item.get("length_type") and length_type_idx >= 0:
                item["length_type"] = self._safe_row_get(row, length_type_idx)
            if not item.get("length_min") and length_min_idx >= 0:
                item["length_min"] = self._safe_row_get(row, length_min_idx)
            if not item.get("length_max") and length_max_idx >= 0:
                item["length_max"] = self._safe_row_get(row, length_max_idx)
            if not item.get("data_type") and datatype_idx >= 0:
                item["data_type"] = self._safe_row_get(row, datatype_idx)
            if not item.get("resolution") and resolution_idx >= 0:
                item["resolution"] = self._safe_row_get(row, resolution_idx)
            if not item.get("offset") and offset_idx >= 0:
                item["offset"] = self._safe_row_get(row, offset_idx)
            if not item.get("unit") and unit_idx >= 0:
                item["unit"] = self._safe_row_get(row, unit_idx)
            if not item.get("discrete_values") and discrete_idx >= 0:
                item["discrete_values"] = self._safe_row_get(row, discrete_idx)
            if not item.get("comment") and comment_idx >= 0:
                item["comment"] = self._safe_row_get(row, comment_idx)

            member_name = self._safe_row_get(row, member_name_idx) if member_name_idx >= 0 else ""
            member_reference = self._safe_row_get(row, member_reference_idx) if member_reference_idx >= 0 else ""
            member_primitive_type = self._safe_row_get(row, datatype_idx) if datatype_idx >= 0 else ""
            if member_name or member_reference or member_primitive_type:
                item["members"].append({
                    "name": member_name,
                    "description": self._safe_row_get(row, member_description_idx) if member_description_idx >= 0 else "",
                    "reference": member_reference,
                    "data_type": member_primitive_type,
                    "resolution": self._safe_row_get(row, resolution_idx) if resolution_idx >= 0 else "",
                    "offset": self._safe_row_get(row, offset_idx) if offset_idx >= 0 else "",
                    "unit": self._safe_row_get(row, unit_idx) if unit_idx >= 0 else "",
                    "discrete_values": self._safe_row_get(row, discrete_idx) if discrete_idx >= 0 else "",
                    "comment": self._safe_row_get(row, comment_idx) if comment_idx >= 0 else "",
                    "initial_value": self._safe_row_get(row, initial_value_idx) if initial_value_idx >= 0 else "",
                    "invalid_value": self._safe_row_get(row, invalid_value_idx) if invalid_value_idx >= 0 else "",
                })

        return registry

    def _parse_serialization_sheet(self, sheet) -> Dict[str, Any]:
        config = self._default_serialization_config()
        rows = list(sheet.iter_rows(values_only=True))
        for row in rows:
            key = self._normalize_header(self._safe_row_get(row, 0))
            value = self._safe_row_get(row, 1)
            if not key:
                continue
            if key == "length field for struct":
                config["struct_has_length_field"] = self._parse_yes_no(value)
            elif key == "struct length field size":
                config["struct_length_field_size_bits"] = self._parse_int(value) or config["struct_length_field_size_bits"]
            elif key == "string length field size":
                config["string_length_field_size_bits"] = self._parse_int(value) or config["string_length_field_size_bits"]
            elif key == "array length field size":
                config["array_length_field_size_bits"] = self._parse_int(value) or config["array_length_field_size_bits"]
            elif key == "string encoding":
                config["string_encoding"] = value or config["string_encoding"]
            elif key == "alignment":
                config["alignment_bits"] = self._parse_int(value) or config["alignment_bits"]
        return config

    def _build_message_schemas(
        self,
        service_entries: List[Dict[str, Any]],
        datatype_registry: Dict[str, Dict[str, Any]],
        serialization_config: Dict[str, Any],
    ) -> Dict[str, List[Dict[str, Any]]]:
        rules: Dict[str, List[Dict[str, Any]]] = {}
        for entry in service_entries:
            rule = {
                "sheet_name": self.SERVICE_INTERFACES_SHEET_NAME,
                "service_id": entry.get("service_id", ""),
                "method_id": entry.get("method_id", ""),
                "method_name": entry.get("method_name", ""),
                "field_name": entry.get("parameter_name", ""),
                "parameter_name": entry.get("parameter_name", ""),
                "parameter_direction": entry.get("direction", ""),
                "parameter_description": entry.get("parameter_description", ""),
                "data_type": entry.get("parameter_data_type", ""),
                "message_kind": entry.get("message_kind", ""),
                "service_interface_name": entry.get("service_interface_name", ""),
                "schema": self._resolve_type_schema(
                    entry.get("parameter_data_type", ""),
                    datatype_registry,
                    serialization_config,
                    visited=None,
                ),
                "_row": int(entry.get("row", 0) or 0),
            }

            base_key = f"{rule['service_id']}:{rule['method_id']}"
            rules.setdefault(base_key, []).append(rule)
            message_kind = str(rule.get("message_kind") or "").strip().lower()
            if message_kind:
                rules.setdefault(f"{base_key}:{message_kind}", []).append(rule)

        for key, items in rules.items():
            items.sort(key=lambda item: (item.get("_row", 0), item.get("field_name", "")))
            for item in items:
                item.pop("_row", None)
        return rules

    def _resolve_type_schema(
        self,
        type_name: Any,
        datatype_registry: Dict[str, Dict[str, Any]],
        serialization_config: Dict[str, Any],
        visited: Optional[Tuple[str, ...]] = None,
    ) -> Dict[str, Any]:
        normalized_name = str(type_name or "").strip()
        if not normalized_name:
            return self._build_unknown_schema("")

        visited = visited or tuple()
        if normalized_name in visited:
            return {
                "category": "reference",
                "type_name": normalized_name,
                "data_type": normalized_name,
                "description": "recursive reference",
            }

        item = datatype_registry.get(normalized_name)
        if not item:
            return self._build_inline_primitive_schema(
                type_name=normalized_name,
                data_type=normalized_name,
                description="",
                resolution="",
                offset="",
                unit="",
                discrete_values="",
                comment="",
                invalid_value="",
                initial_value="",
                data_category="",
            )

        data_category = self._normalize_header(item.get("data_category"))
        next_visited = visited + (normalized_name,)

        if data_category == "struct":
            return {
                "category": "struct",
                "type_name": normalized_name,
                "description": item.get("description", ""),
                "has_length_prefix": bool(serialization_config.get("struct_has_length_field", False)),
                "length_field_size_bits": int(serialization_config.get("struct_length_field_size_bits", 32) or 32),
                "members": [
                    {
                        "name": member.get("name", "") or member.get("reference", "") or member.get("data_type", "") or "member",
                        "description": member.get("description", ""),
                        "type_name": member.get("reference") or member.get("data_type", ""),
                        "schema": self._resolve_member_schema(member, datatype_registry, serialization_config, next_visited),
                    }
                    for member in item.get("members", [])
                ],
            }

        if data_category == "array":
            element_schema = self._resolve_member_schema({
                "name": item.get("members", [{}])[0].get("name", "") if item.get("members") else "",
                "reference": item.get("members", [{}])[0].get("reference", "") if item.get("members") else "",
                "data_type": item.get("members", [{}])[0].get("data_type", item.get("data_type", "")) if item.get("members") else item.get("data_type", ""),
                "description": item.get("members", [{}])[0].get("description", "") if item.get("members") else item.get("description", ""),
                "resolution": item.get("members", [{}])[0].get("resolution", item.get("resolution", "")) if item.get("members") else item.get("resolution", ""),
                "offset": item.get("members", [{}])[0].get("offset", item.get("offset", "")) if item.get("members") else item.get("offset", ""),
                "unit": item.get("members", [{}])[0].get("unit", item.get("unit", "")) if item.get("members") else item.get("unit", ""),
                "discrete_values": item.get("members", [{}])[0].get("discrete_values", item.get("discrete_values", "")) if item.get("members") else item.get("discrete_values", ""),
                "comment": item.get("members", [{}])[0].get("comment", item.get("comment", "")) if item.get("members") else item.get("comment", ""),
                "invalid_value": item.get("members", [{}])[0].get("invalid_value", item.get("invalid_value", "")) if item.get("members") else item.get("invalid_value", ""),
                "initial_value": item.get("members", [{}])[0].get("initial_value", item.get("initial_value", "")) if item.get("members") else item.get("initial_value", ""),
            }, datatype_registry, serialization_config, next_visited)
            return {
                "category": "array",
                "type_name": normalized_name,
                "description": item.get("description", ""),
                "length_type": self._normalize_length_type(item.get("length_type", "")),
                "min_length": self._parse_int(item.get("length_min")),
                "max_length": self._parse_int(item.get("length_max")),
                "length_field_size_bits": int(serialization_config.get("array_length_field_size_bits", 32) or 32),
                "element_name": item.get("members", [{}])[0].get("name", "item") if item.get("members") else "item",
                "element_schema": element_schema,
            }

        if data_category == "string":
            return {
                "category": "string",
                "type_name": normalized_name,
                "description": item.get("description", ""),
                "encoding": item.get("data_type") or serialization_config.get("string_encoding", "utf-8"),
                "length_type": self._normalize_length_type(item.get("length_type", "")),
                "min_length": self._parse_int(item.get("length_min")),
                "max_length": self._parse_int(item.get("length_max")),
                "length_field_size_bits": int(serialization_config.get("string_length_field_size_bits", 32) or 32),
            }

        return self._build_inline_primitive_schema(
            type_name=normalized_name,
            data_type=item.get("data_type", normalized_name),
            description=item.get("description", ""),
            resolution=item.get("resolution", ""),
            offset=item.get("offset", ""),
            unit=item.get("unit", ""),
            discrete_values=item.get("discrete_values", ""),
            comment=item.get("comment", ""),
            invalid_value=item.get("invalid_value", ""),
            initial_value=item.get("initial_value", ""),
            data_category=item.get("data_category", ""),
        )

    def _resolve_member_schema(
        self,
        member: Dict[str, Any],
        datatype_registry: Dict[str, Dict[str, Any]],
        serialization_config: Dict[str, Any],
        visited: Tuple[str, ...],
    ) -> Dict[str, Any]:
        reference = str(member.get("reference") or "").strip()
        if reference:
            return self._resolve_type_schema(reference, datatype_registry, serialization_config, visited)
        data_type = str(member.get("data_type") or "").strip()
        if data_type:
            if data_type in datatype_registry:
                return self._resolve_type_schema(data_type, datatype_registry, serialization_config, visited)
            return self._build_inline_primitive_schema(
                type_name=data_type,
                data_type=data_type,
                description=member.get("description", ""),
                resolution=member.get("resolution", ""),
                offset=member.get("offset", ""),
                unit=member.get("unit", ""),
                discrete_values=member.get("discrete_values", ""),
                comment=member.get("comment", ""),
                invalid_value=member.get("invalid_value", ""),
                initial_value=member.get("initial_value", ""),
                data_category="",
            )
        return self._build_unknown_schema(str(member.get("name") or "member"))

    def _build_inline_primitive_schema(
        self,
        type_name: str,
        data_type: str,
        description: str,
        resolution: str,
        offset: str,
        unit: str,
        discrete_values: str,
        comment: str,
        invalid_value: str,
        initial_value: str,
        data_category: str,
    ) -> Dict[str, Any]:
        return {
            "category": self._normalize_primitive_category(data_category, discrete_values),
            "type_name": str(type_name or data_type or ""),
            "data_type": str(data_type or type_name or ""),
            "description": str(description or ""),
            "resolution": str(resolution or ""),
            "offset": str(offset or ""),
            "unit": str(unit or ""),
            "comment": str(comment or ""),
            "invalid_value": str(invalid_value or ""),
            "initial_value": str(initial_value or ""),
            "enum_map": self._parse_discrete_values(discrete_values),
        }

    @staticmethod
    def _build_unknown_schema(type_name: str) -> Dict[str, Any]:
        return {
            "category": "unknown",
            "type_name": str(type_name or ""),
            "data_type": str(type_name or ""),
            "description": "",
        }

    @staticmethod
    def _default_serialization_config() -> Dict[str, Any]:
        return {
            "alignment_bits": 8,
            "struct_has_length_field": True,
            "struct_length_field_size_bits": 32,
            "string_length_field_size_bits": 32,
            "array_length_field_size_bits": 32,
            "string_encoding": "UTF-8",
        }

    def _find_service_interface_header_row_index(self, rows: List[Tuple[Any, ...]], max_scan: int = 10) -> int:
        best_index = 0
        best_score = -1
        for idx, row in enumerate(rows[:max_scan]):
            headers = [self._normalize_header(cell) for cell in row]
            score = 0
            if self._find_col_idx(headers, ["service interface name"]) >= 0:
                score += 2
            if self._find_col_idx(headers, ["service id"]) >= 0:
                score += 2
            if self._find_col_idx(headers, ["method/event/field"]) >= 0:
                score += 2
            if self._find_col_idx(headers, ["parameter name"]) >= 0:
                score += 2
            if self._find_col_idx(headers, ["parameter data type"]) >= 0:
                score += 2
            if any(headers):
                score += 1
            if score > best_score:
                best_score = score
                best_index = idx
        return best_index

    def _find_datatype_header_row_index(self, rows: List[Tuple[Any, ...]], max_scan: int = 10) -> int:
        best_index = 0
        best_score = -1
        for idx, row in enumerate(rows[:max_scan]):
            headers = [self._normalize_header(cell) for cell in row]
            score = 0
            if self._find_col_idx(headers, ["parameter data type name"]) >= 0:
                score += 2
            if self._find_col_idx(headers, ["data category"]) >= 0:
                score += 2
            if self._find_col_idx(headers, ["member datatype reference"]) >= 0:
                score += 2
            if self._find_col_idx(headers, ["datatype"]) >= 0:
                score += 2
            if any(headers):
                score += 1
            if score > best_score:
                best_score = score
                best_index = idx
        return best_index

    @staticmethod
    def _normalize_message_kind(message_type: Optional[Any]) -> str:
        if message_type is None or message_type == "":
            return ""
        if isinstance(message_type, str):
            normalized = message_type.strip().lower()
            if normalized in {"request", "response", "notification"}:
                return normalized
            try:
                value = int(normalized, 16) if normalized.startswith("0x") else int(normalized, 10)
            except ValueError:
                return ""
        else:
            try:
                value = int(message_type)
            except (TypeError, ValueError):
                return ""

        base_type = value & 0xDF
        if base_type in {0x00, 0x01}:
            return "request"
        if base_type == 0x02:
            return "notification"
        if base_type in {0x80, 0x81}:
            return "response"
        return ""

    def _resolve_message_kind_from_row(self, method_kind: str, accessor: str, direction: str) -> str:
        normalized_method_kind = self._normalize_header(method_kind)
        normalized_accessor = self._normalize_header(accessor)
        normalized_direction = str(direction or "").strip().upper()

        if "event" in normalized_method_kind or normalized_accessor == "notifier":
            return "notification"
        if "field" in normalized_method_kind and normalized_accessor == "notifier":
            return "notification"
        if "field" in normalized_method_kind and normalized_accessor == "getter":
            return "response"
        if "field" in normalized_method_kind and normalized_accessor == "setter":
            return "request"
        if normalized_direction == "OUT" and "field" in normalized_method_kind and normalized_accessor == "getter":
            return "response"
        if normalized_direction == "IN":
            return "request"
        if normalized_direction == "OUT":
            return "response"
        return ""

    @staticmethod
    def _normalize_length_type(value: Any) -> str:
        normalized = str(value or "").strip().lower()
        if "dynamic" in normalized:
            return "dynamic"
        if "fixed" in normalized:
            return "fixed"
        return normalized

    @staticmethod
    def _normalize_primitive_category(data_category: Any, discrete_values: Any) -> str:
        normalized_category = str(data_category or "").strip().lower()
        if "enum" in normalized_category or str(discrete_values or "").strip():
            return "enum"
        if "bool" in normalized_category:
            return "bool"
        return "primitive"

    @staticmethod
    def _parse_yes_no(value: Any) -> bool:
        return str(value or "").strip().lower() in {"yes", "true", "1", "y"}

    @classmethod
    def _parse_discrete_values(cls, value: Any) -> Dict[int, str]:
        text = str(value or "").strip()
        if not text:
            return {}
        mapping: Dict[int, str] = {}
        
        clean_text = text.replace('\n', ' ').replace('\r', ' ').replace(';', ' ')
        # Find all occurrences of something like "0x1=" or "1="
        import re
        parts = re.split(r'((?:0x[0-9a-fA-F]+|\d+)\s*=)', clean_text)
        
        # parts will look like: 
        # ['', '0x0=', 'Ok, ', '0x1=', 'NOK']
        
        current_key = None
        for chunk in parts:
            if not chunk:
                continue
                
            chunk_stripped = chunk.strip()
            if chunk_stripped.endswith('='):
                # This is a key
                key_str = chunk_stripped[:-1].strip()
                try:
                    current_key = int(key_str, 16) if key_str.lower().startswith("0x") else int(key_str, 10)
                except ValueError:
                    current_key = None
            else:
                # This is a value
                if current_key is not None:
                    val_str = chunk.strip(" ,;\t")
                    mapping[current_key] = val_str
                    current_key = None
                    
        return mapping

    def _parse_deployment_sheet(self, sheet) -> Dict[str, Dict[str, Any]]:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return {}

        header_row_index = self._find_header_row_index(rows)
        headers = [self._normalize_header(v) for v in rows[header_row_index]]

        service_id_idx = self._find_col_idx(headers, [
            "service id", "serviceid", "service_identifier", "service identifier"
        ])
        instance_id_idx = self._find_col_idx(headers, [
            "instance id", "instanceid", "instance identifier", "instance_identifier"
        ])
        interface_name_idx = self._find_col_idx(headers, [
            "service interface name", "serviceinterface name", "service interface", "interface name"
        ])

        missing = []
        if service_id_idx < 0:
            missing.append("Service ID")
        if instance_id_idx < 0:
            missing.append("Instance ID")
        if interface_name_idx < 0:
            missing.append("Service Interface Name")
        if missing:
            raise SomeIPMatrixError(f"Deployment sheet 缺少关键列: {', '.join(missing)}")

        services: Dict[str, Dict[str, Any]] = {}
        for row in rows[header_row_index + 1:]:
            service_id_raw = self._safe_row_get(row, service_id_idx)
            instance_id_raw = self._safe_row_get(row, instance_id_idx)
            interface_name = self._safe_row_get(row, interface_name_idx)

            normalized_service_id = self._normalize_id(service_id_raw)
            normalized_instance_id = self._normalize_id(instance_id_raw)
            if not normalized_service_id or not normalized_instance_id or not interface_name:
                continue

            lookup_key = self._build_lookup_key(normalized_service_id, normalized_instance_id)
            services[lookup_key] = {
                "service_id": normalized_service_id,
                "instance_id": normalized_instance_id,
                "service_interface_name": interface_name,
            }

        return services

    def _parse_payload_sheets(self, workbook, services: Dict[str, Dict[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
        service_name_index = self._build_service_name_index(services)
        rules: Dict[str, List[Dict[str, Any]]] = {}

        for sheet_name in workbook.sheetnames:
            if sheet_name == self.DEPLOYMENT_SHEET_NAME:
                continue
            sheet_rules = self._parse_payload_sheet(workbook[sheet_name], sheet_name, service_name_index)
            for key, items in sheet_rules.items():
                rules.setdefault(key, []).extend(items)

        for key, items in rules.items():
            items.sort(key=lambda item: (item.get("byte_offset", 0), item.get("bit_offset", 0), item.get("field_name", "")))
        return rules

    def _parse_payload_sheet(self, sheet, sheet_name: str, service_name_index: Dict[str, str]) -> Dict[str, List[Dict[str, Any]]]:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return {}

        header_row_index = self._find_payload_header_row_index(rows)
        headers = [self._normalize_header(v) for v in rows[header_row_index]]

        field_name_idx = self._find_col_idx(headers, [
            "field name", "signal name", "parameter name", "param name", "data name", "name"
        ])
        method_id_idx = self._find_col_idx(headers, [
            "method id", "event id", "message id"
        ])
        service_id_idx = self._find_col_idx(headers, [
            "service id", "serviceid"
        ])
        service_name_idx_col = self._find_col_idx(headers, [
            "service interface name", "service name", "interface name"
        ])
        method_name_idx = self._find_col_idx(headers, [
            "method name", "event name", "message name", "api name"
        ])
        byte_offset_idx = self._find_col_idx(headers, [
            "byte offset", "start byte", "byte position", "bytepos", "offset byte", "payload offset"
        ])
        bit_offset_idx = self._find_col_idx(headers, [
            "bit offset", "start bit", "bit position", "bitpos"
        ])
        bit_length_idx = self._find_col_idx(headers, [
            "bit length", "signal length", "length(bit)", "len(bit)"
        ])
        byte_length_idx = self._find_col_idx(headers, [
            "byte length", "length(byte)", "len(byte)"
        ])
        type_idx = self._find_col_idx(headers, ["data type", "datatype", "type"])
        direction_idx = self._find_col_idx(headers, [
            "direction", "request/response", "request response", "req/rsp", "message type"
        ])
        unit_idx = self._find_col_idx(headers, ["unit"])
        factor_idx = self._find_col_idx(headers, ["factor", "scale"])
        offset_idx = self._find_col_idx(headers, ["offset", "bias"])

        if field_name_idx < 0 or method_id_idx < 0:
            return {}
        if byte_offset_idx < 0 and bit_offset_idx < 0:
            return {}

        inferred_service_id = self._infer_service_id_from_name(sheet_name, service_name_index)
        last_service_id = inferred_service_id
        last_method_id = ""
        last_method_name = ""
        parsed: Dict[str, List[Dict[str, Any]]] = {}

        for row in rows[header_row_index + 1:]:
            field_name = self._safe_row_get(row, field_name_idx)
            if not field_name:
                continue

            service_id = self._normalize_id(self._safe_row_get(row, service_id_idx)) if service_id_idx >= 0 else ""
            if not service_id and service_name_idx_col >= 0:
                service_name = self._safe_row_get(row, service_name_idx_col)
                service_id = self._infer_service_id_from_name(service_name, service_name_index)
            if not service_id:
                service_id = last_service_id
            if service_id:
                last_service_id = service_id

            method_id = self._normalize_id(self._safe_row_get(row, method_id_idx)) if method_id_idx >= 0 else ""
            if not method_id:
                method_id = last_method_id
            if method_id:
                last_method_id = method_id

            method_name = self._safe_row_get(row, method_name_idx) if method_name_idx >= 0 else ""
            if method_name:
                last_method_name = method_name
            else:
                method_name = last_method_name

            if not service_id or not method_id:
                continue

            byte_offset = self._parse_int(self._safe_row_get(row, byte_offset_idx)) if byte_offset_idx >= 0 else 0
            bit_offset = self._parse_int(self._safe_row_get(row, bit_offset_idx)) if bit_offset_idx >= 0 else 0
            bit_length = self._parse_int(self._safe_row_get(row, bit_length_idx)) if bit_length_idx >= 0 else 0
            if bit_length <= 0 and byte_length_idx >= 0:
                byte_length = self._parse_int(self._safe_row_get(row, byte_length_idx))
                bit_length = byte_length * 8 if byte_length > 0 else 0
            if bit_length <= 0:
                bit_length = 8

            key = f"{service_id}:{method_id}"
            parsed.setdefault(key, []).append({
                "sheet_name": sheet_name,
                "service_id": service_id,
                "method_id": method_id,
                "method_name": method_name,
                "field_name": field_name,
                "byte_offset": max(0, byte_offset),
                "bit_offset": max(0, bit_offset),
                "bit_length": max(1, bit_length),
                "data_type": self._safe_row_get(row, type_idx) if type_idx >= 0 else "",
                "direction": self._safe_row_get(row, direction_idx) if direction_idx >= 0 else "",
                "unit": self._safe_row_get(row, unit_idx) if unit_idx >= 0 else "",
                "factor": self._safe_row_get(row, factor_idx) if factor_idx >= 0 else "",
                "offset_value": self._safe_row_get(row, offset_idx) if offset_idx >= 0 else "",
            })

        return parsed

    def _find_header_row_index(self, rows: List[Tuple[Any, ...]], max_scan: int = 20) -> int:
        best_index = 0
        best_score = -1
        for idx, row in enumerate(rows[:max_scan]):
            headers = [self._normalize_header(cell) for cell in row]
            score = 0
            if self._find_col_idx(headers, ["service id", "serviceid"]) >= 0:
                score += 2
            if self._find_col_idx(headers, ["instance id", "instanceid"]) >= 0:
                score += 2
            if self._find_col_idx(headers, ["service interface name", "interface name"]) >= 0:
                score += 2
            if any(headers):
                score += 1
            if score > best_score:
                best_score = score
                best_index = idx
        return best_index

    def _find_payload_header_row_index(self, rows: List[Tuple[Any, ...]]) -> int:
        best_index = 0
        best_score = -1
        for idx, row in enumerate(rows[:self.PAYLOAD_HEADER_SCAN_LIMIT]):
            headers = [self._normalize_header(cell) for cell in row]
            score = 0
            if self._find_col_idx(headers, ["service id", "service name", "interface name"]) >= 0:
                score += 2
            if self._find_col_idx(headers, ["method id", "event id", "message id"]) >= 0:
                score += 2
            if self._find_col_idx(headers, ["field name", "signal name", "parameter name", "name"]) >= 0:
                score += 3
            if self._find_col_idx(headers, ["byte offset", "start byte", "bit offset", "start bit"]) >= 0:
                score += 3
            if any(headers):
                score += 1
            if score > best_score:
                best_score = score
                best_index = idx
        return best_index

    @staticmethod
    def _build_service_name_index(services: Dict[str, Dict[str, Any]]) -> Dict[str, str]:
        index: Dict[str, str] = {}
        for item in services.values():
            service_name = str(item.get("service_interface_name") or "").strip().lower()
            service_id = str(item.get("service_id") or "").strip()
            if service_name and service_id and service_name not in index:
                index[service_name] = service_id
        return index

    @staticmethod
    def _infer_service_id_from_name(name: str, service_name_index: Dict[str, str]) -> str:
        normalized = str(name or "").strip().lower()
        if not normalized:
            return ""
        if normalized in service_name_index:
            return service_name_index[normalized]
        for service_name, service_id in service_name_index.items():
            if service_name and service_name in normalized:
                return service_id
        return ""

    @staticmethod
    def _find_col_idx(headers: Iterable[str], keywords: Iterable[str]) -> int:
        for idx, header in enumerate(headers):
            if any(keyword in header for keyword in keywords):
                return idx
        return -1

    @staticmethod
    def _normalize_header(value: Any) -> str:
        if value is None:
            return ""
        return " ".join(str(value).replace("\n", " ").replace("\t", " ").split()).strip().lower()

    @staticmethod
    def _safe_row_get(row: Tuple[Any, ...], idx: int) -> str:
        if idx < 0 or idx >= len(row):
            return ""
        value = row[idx]
        return "" if value is None else str(value).strip()

    @classmethod
    def _build_lookup_key(cls, service_id: Any, instance_id: Any) -> str:
        normalized_service_id = cls._normalize_id(service_id)
        normalized_instance_id = cls._normalize_id(instance_id)
        return f"{normalized_service_id}:{normalized_instance_id}"

    @staticmethod
    def _normalize_id(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, int):
            return f"0x{value:04X}"
        if isinstance(value, float):
            if value.is_integer():
                return f"0x{int(value):04X}"
            return ""

        text = str(value).strip()
        if not text:
            return ""
        compact = text.replace("_", "").replace(" ", "")
        try:
            if compact.lower().startswith("0x"):
                return f"0x{int(compact, 16):04X}"
            return f"0x{int(compact, 10):04X}"
        except ValueError:
            return text.upper()

    @staticmethod
    def _parse_int(value: Any) -> int:
        text = str(value or "").strip()
        if not text:
            return 0
        match = re.search(r"0x[0-9A-Fa-f]+|\d+", text)
        if not match:
            return 0
        token = match.group(0)
        return int(token, 16) if token.lower().startswith("0x") else int(token, 10)

    @staticmethod
    def _build_fingerprint(file_path: Path) -> Dict[str, Any]:
        stat = file_path.stat()
        return {
            "mtime": stat.st_mtime,
            "size": stat.st_size,
        }

    def _cache_file_path(self) -> Path:
        return Path(get_device_config_path()) / f"{self.CACHE_NAME}.json"

    def _load_disk_cache(self) -> Dict[str, Any]:
        data = self._config_service.load_json(self.CACHE_NAME, default={})
        return data if isinstance(data, dict) else {}

    def _save_disk_cache(self, data: Dict[str, Any]) -> None:
        self._config_service.save_json(self.CACHE_NAME, data)

    def _load_app_settings(self) -> Dict[str, Any]:
        data = self._config_service.load_json("app_settings", default={})
        return data if isinstance(data, dict) else {}


_default_store = SomeIPMatrixStore()


def import_someip_matrix(file_path: Optional[str] = None, force_reload: bool = False) -> Dict[str, Any]:
    return _default_store.import_matrix(file_path=file_path, force_reload=force_reload)


def get_service_interface_name(service_id: Any, instance_id: Any, file_path: Optional[str] = None, default: str = "Unknown Service") -> str:
    return _default_store.get_service_interface_name(service_id, instance_id, file_path=file_path, default=default)


def get_service_info(service_id: Any, instance_id: Any, file_path: Optional[str] = None) -> Optional[Dict[str, Any]]:
    return _default_store.get_service_info(service_id, instance_id, file_path=file_path)


def get_payload_rules(
    service_id: Any,
    method_id: Any,
    file_path: Optional[str] = None,
    message_type: Optional[Any] = None,
) -> List[Dict[str, Any]]:
    return _default_store.get_payload_rules(
        service_id,
        method_id,
        file_path=file_path,
        message_type=message_type,
    )

