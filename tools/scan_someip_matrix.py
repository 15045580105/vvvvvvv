import json
from pathlib import Path

import openpyxl

MATRIX_PATH = Path(r"/app/common/file/GAC_VehicleA_SOMEIP_CMX_CCU_Phase1_(CCU_S32G_Acore)_V2.2.38_20251216.xlsx")
OUTPUT_PATH = Path(r"D:\VCC\tools\matrix_sheet_scan.json")


def main() -> None:
    workbook = openpyxl.load_workbook(str(MATRIX_PATH), read_only=True, data_only=True)
    try:
        data = {}
        for name in workbook.sheetnames:
            sheet = workbook[name]
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [str(value).strip() if value is not None else "" for value in row]
                if any(values):
                    rows.append({"row": index, "values": values[:30]})
                if len(rows) >= 6:
                    break
            data[name] = rows
        OUTPUT_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        print(str(OUTPUT_PATH))
    finally:
        workbook.close()


if __name__ == "__main__":
    main()

