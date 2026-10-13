#!/usr/bin/env python3
"""提取 ICC SOMEIP 完整服务定义，输出所有 Server/Client 通信矩阵。"""

import openpyxl
import json

EXCEL_PATH = r'd:\奇瑞\ICC SOMEIP Services Interface Description V4.400（For EEA5.1）20251030.xlsx'

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb.worksheets[5]  # "3 SOMEIP通信行为定义"

# 解析所有数据行 (row 2 ~ max_row)
services = []
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
    vals = [c.value for c in row]
    if vals[0] is None and vals[1] is None:
        continue  # 空行

    svc = {
        "service_name": str(vals[0] or "").strip(),
        "service_id": str(vals[1] or "").strip(),
        "instance_id": vals[2],
        "major_version": vals[3],
        "minor_version": vals[4],
        "ttl": vals[5],
        "vlan_id": vals[6],
        "vlan_name": str(vals[7] or "").strip(),
        "vlan_priority": vals[8],
        "transport_protocol": str(vals[9] or "").strip(),
        "server_ecu": str(vals[10] or "").strip(),
        "server_ip": str(vals[11] or "").strip(),
        "server_subnet": str(vals[12] or "").strip(),
        "server_port": vals[13],
        "client_ecu": str(vals[14] or "").strip(),
        "client_id": str(vals[15] or "").strip(),
        "client_ip": str(vals[16] or "").strip(),
        "client_subnet": str(vals[17] or "").strip(),
        "client_port": vals[18],
        "remark": str(vals[19] or "").strip() if len(vals) > 19 else "",
    }
    services.append(svc)

# 按 Server 分组输出
print(f"共 {len(services)} 条服务定义\n")

# 按 server_ecu 分组
from collections import defaultdict
by_server = defaultdict(list)
for s in services:
    by_server[s["server_ecu"]].append(s)

print("=" * 100)
print("  按 Server ECU 分组的服务列表")
print("=" * 100)

for server_ecu in sorted(by_server.keys()):
    svcs = by_server[server_ecu]
    print(f"\n--- Server: {server_ecu} ({len(svcs)} 条) ---")
    for s in svcs:
        svc_id_hex = s["service_id"]
        print(f"  {s['service_name']:30s} | SvcID=0x{svc_id_hex:>6s} | Inst={s['instance_id']} "
              f"| v{s['major_version']}.{s['minor_version']} | TTL={s['ttl']} "
              f"| {s['transport_protocol']:10s} | VLAN={s['vlan_id']}")
        print(f"    Server: {s['server_ecu']:20s} {s['server_ip']:>16s}:{s['server_port']}")
        print(f"    Client: {s['client_ecu']:20s} {s['client_ip']:>16s}:{s['client_port']} (ClientID={s['client_id']})")

# 特别输出 ICC 作为 Server 的完整列表（这是我们作为 Client 要测试的）
print("\n\n" + "=" * 100)
print("  ICC 作为 Server 的服务 (VCC 作为 Client 测试)")
print("=" * 100)

icc_services = [s for s in services if "ICC" in s["server_ecu"]]
for s in icc_services:
    print(f"\n  {s['service_name']} | 0x{s['service_id']} | {s['transport_protocol']} "
          f"| Server={s['server_ip']}:{s['server_port']} "
          f"| Client={s['client_ecu']}({s['client_ip']}:{s['client_port']})")

# 输出 JSON 方便后续使用
output = []
for s in services:
    output.append({
        "service_name": s["service_name"],
        "service_id": s["service_id"],
        "instance_id": s["instance_id"],
        "major_version": s["major_version"],
        "minor_version": s["minor_version"],
        "ttl": s["ttl"],
        "vlan_id": s["vlan_id"],
        "transport_protocol": s["transport_protocol"],
        "server_ecu": s["server_ecu"],
        "server_ip": s["server_ip"],
        "server_port": s["server_port"],
        "client_ecu": s["client_ecu"],
        "client_ip": s["client_ip"],
        "client_port": s["client_port"],
    })

with open(r"d:\VCC\tools\icc_service_defs.json", "w", encoding="utf-8") as f:
    json.dump(output, f, ensure_ascii=False, indent=2)
print(f"\n\n已导出完整定义到 tools/icc_service_defs.json")
#!/usr/bin/env python3
"""解析 ICC SOMEIP Excel 接口文档，提取完整的 Client/Server 服务定义。"""

import openpyxl
import json

EXCEL_PATH = r'd:\奇瑞\ICC SOMEIP Services Interface Description V4.400（For EEA5.1）20251030.xlsx'

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb.worksheets[5]  # "3 SOMEIP通信行为定义"
print(f"Sheet: {ws.title}")
print(f"Dimensions: {ws.dimensions}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

# 先打印前20行，了解表头结构
print("\n=== 前 20 行内容 ===")
for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=False), 1):
    vals = []
    for cell in row:
        v = cell.value
        if v is not None:
            vals.append(f"[{cell.column}]{v}")
    if vals:
        print(f"  Row {row_idx}: {' | '.join(vals)}")

# 找到表头行（包含 "Service ID" 或类似关键字）
print("\n=== 查找表头行 ===")
header_row = None
for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=False), 1):
    vals = [str(c.value).strip().lower() if c.value else "" for c in row]
    combined = " ".join(vals)
    if "service" in combined and ("id" in combined or "name" in combined):
        header_row = row_idx
        print(f"  表头行: {row_idx}")
        for cell in row:
            if cell.value is not None:
                print(f"    Col {cell.column}: {cell.value}")
        break

if header_row is None:
    # 尝试找含 "server" 或 "client" 的行
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=False), 1):
        vals = [str(c.value).strip().lower() if c.value else "" for c in row]
        combined = " ".join(vals)
        if "server" in combined or "client" in combined:
            print(f"  候选行 {row_idx}: {combined[:200]}")
