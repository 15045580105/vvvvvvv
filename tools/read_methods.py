#!/usr/bin/env python3
"""Read method IDs from Excel matrix."""
import sys
sys.path.insert(0, r"d:\VCC")
import openpyxl

EXCEL = r"d:\奇瑞\ICC SOMEIP Services Interface Description V4.400（For EEA5.1）20251030.xlsx"
wb = openpyxl.load_workbook(EXCEL, read_only=True, data_only=True)

# Find the service definition sheet
target = None
for name in wb.sheetnames:
    sheet = wb[name]
    for row in sheet.iter_rows(values_only=True, max_row=5):
        vals = [str(v).strip().lower() if v else "" for v in row]
        joined = " ".join(vals)
        if "service id" in joined and "method" in joined:
            target = name
            break
    if target:
        break

if not target:
    print("ERROR: service definition sheet not found")
    sys.exit(1)

print(f"Sheet: {target!r}\n")
sheet = wb[target]
rows = list(sheet.iter_rows(values_only=True))

# Find header
header_idx = 0
for idx, row in enumerate(rows):
    vals = [str(v).strip().lower() if v else "" for v in row]
    joined = " ".join(vals)
    if "service id" in joined and "method" in joined:
        header_idx = idx
        headers = [str(v).strip() if v else "" for v in row]
        print(f"Headers: {headers}")
        break

# Parse data
print(f"\n{'Service Name':<30} {'SvcID':>6} {'Type':<10} {'Method/Event':<30} {'MethodID':>8} {'EventGrp':>10} {'Proto':>6}")
print("-" * 110)

services = {}
for row in rows[header_idx+1:]:
    vals = [str(v).strip() if v else "" for v in row]
    if not vals[0] and not vals[1]:
        continue
    svc_name = vals[0]
    try:
        svc_id = int(vals[1], 0) if vals[1] else 0
    except:
        continue
    rpc_type = vals[3] if len(vals) > 3 else ""  # col D
    method_name = vals[7] if len(vals) > 7 else ""  # col H
    try:
        method_id = int(vals[8], 0) if len(vals) > 8 and vals[8] else 0  # col I
    except:
        method_id = 0
    eventgroup = vals[10] if len(vals) > 10 else ""  # col K
    proto = vals[11] if len(vals) > 11 else ""  # col L

    print(f"  {svc_name:<28} 0x{svc_id:04X} {rpc_type:<10} {method_name:<28} 0x{method_id:04X} {eventgroup:>10} {proto:>6}")

    if svc_id not in services:
        services[svc_id] = {"name": svc_name, "methods": [], "events": []}
    if "event" in rpc_type.lower():
        services[svc_id]["events"].append((method_name, method_id, eventgroup))
    else:
        services[svc_id]["methods"].append((method_name, method_id))

print(f"\n\n=== Summary by Service ===")
for sid in sorted(services):
    info = services[sid]
    print(f"\n0x{sid:04X} ({info['name']}):")
    for m_name, m_id in info["methods"]:
        print(f"  RPC: 0x{m_id:04X} - {m_name}")
    for e_name, e_id, eg in info["events"]:
        print(f"  EVT: 0x{e_id:04X} - {e_name} (Group: {eg})")
