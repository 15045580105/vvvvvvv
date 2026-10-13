#!/usr/bin/env python3
"""Full Excel matrix dump: deployment + service definitions."""
import sys
sys.path.insert(0, r"d:\VCC")
import openpyxl

EXCEL = r"d:\奇瑞\ICC SOMEIP Services Interface Description V4.400（For EEA5.1）20251030.xlsx"
wb = openpyxl.load_workbook(EXCEL, read_only=True, data_only=True)
OUT = r"d:\VCC\tools\excel_full_dump.txt"

out = []
def log(m):
    out.append(str(m))

# === Part 1: Deployment sheet ===
log("=" * 80)
log("PART 1: Deployment (服务部署)")
log("=" * 80)

dep_sheet = None
for name in wb.sheetnames:
    sheet = wb[name]
    for row in sheet.iter_rows(values_only=True, max_row=5):
        vals = [str(v).strip().lower() if v else "" for v in row]
        joined = " ".join(vals)
        if "service id" in joined and "server ip" in joined:
            dep_sheet = name
            break
    if dep_sheet:
        break

if dep_sheet:
    log(f"Sheet: {dep_sheet!r}")
    sheet = wb[dep_sheet]
    rows = list(sheet.iter_rows(values_only=True))
    # Find header
    header_idx = 0
    for idx, row in enumerate(rows):
        vals = [str(v).strip().lower() if v else "" for v in row]
        if "service id" in " ".join(vals) and "server ip" in " ".join(vals):
            header_idx = idx
            headers = [str(v).strip() if v else "" for v in row]
            log(f"Headers ({len(headers)} cols): {headers[:15]}")
            break
    # Dump all data rows
    for i, row in enumerate(rows[header_idx+1:]):
        vals = [str(v).strip() if v else "" for v in row[:15]]
        if not any(vals[:3]):
            continue
        log(f"  [{i+1}] {' | '.join(vals[:13])}")
else:
    log("Deployment sheet NOT found")

# === Part 2: Service definition sheet ===
log("\n" + "=" * 80)
log("PART 2: Service Definitions (服务方法/事件定义)")
log("=" * 80)

svc_sheet = None
for name in wb.sheetnames:
    sheet = wb[name]
    for row in sheet.iter_rows(values_only=True, max_row=5):
        vals = [str(v).strip().lower() if v else "" for v in row]
        if "service id" in " ".join(vals) and "method" in " ".join(vals):
            svc_sheet = name
            break
    if svc_sheet:
        break

if svc_sheet:
    log(f"Sheet: {svc_sheet!r}")
    sheet = wb[svc_sheet]
    rows = list(sheet.iter_rows(values_only=True))
    header_idx = 0
    for idx, row in enumerate(rows):
        vals = [str(v).strip().lower() if v else "" for v in row]
        if "service id" in " ".join(vals) and "method" in " ".join(vals):
            header_idx = idx
            headers = [str(v).strip() if v else "" for v in row]
            log(f"Headers ({len(headers)} cols): {headers[:12]}")
            break
    # Dump all data
    current_svc = ""
    current_svc_id = ""
    for i, row in enumerate(rows[header_idx+1:]):
        vals = [str(v).strip() if v else "" for v in row[:12]]
        # Service name/ID only on first row of each service group
        if vals[0]:
            current_svc = vals[0]
            current_svc_id = vals[1]
        rpc_type = vals[3] if len(vals) > 3 else ""
        method_name = vals[7] if len(vals) > 7 else ""
        method_id = vals[8] if len(vals) > 8 else ""
        eventgroup = vals[10] if len(vals) > 10 else ""
        proto = vals[11] if len(vals) > 11 else ""
        if method_name or rpc_type:
            log(f"  {current_svc:<28} [{current_svc_id:>5}] {rpc_type:<8} {method_name:<30} [{method_id:>5}] {eventgroup:<40} {proto}")
else:
    log("Service definition sheet NOT found")

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(out))
print(f"Written to {OUT}")
